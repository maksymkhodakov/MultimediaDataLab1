"""
Допоміжні функції для формування звіту у форматі Word (.docx) за ДСТУ 3008:2015
«Інформація та документація. Звіти у сфері науки і техніки. Структура та правила
оформлення».

Основні вимоги стандарту, які тут реалізовано:
  * аркуш А4, поля: ліве 30 мм, праве 15 мм, верхнє і нижнє 20 мм;
  * шрифт Times New Roman 14 пт, міжрядковий інтервал 1,5, абзацний відступ 1,25 см,
    вирівнювання за шириною;
  * заголовки структурних частин (ЗМІСТ, ВСТУП, ВИСНОВКИ, ПЕРЕЛІК ДЖЕРЕЛ
    ПОСИЛАННЯ) — великими літерами, по центру; кожен розділ — з нової сторінки;
  * номери сторінок — арабськими цифрами у правому верхньому куті; титульний
    аркуш входить до нумерації, але номер на ньому не ставиться;
  * підпис рисунка — під рисунком по центру: «Рисунок 3.1 – Назва»;
  * назва таблиці — над таблицею зліва: «Таблиця 3.1 – Назва»;
  * формули — по центру, номер у круглих дужках біля правого поля.

Документ відкривається з файлу, де вже є титульна сторінка, і доповнюється.
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

FONT = "Times New Roman"
CODE_FONT = "Courier New"
# Ширина області тексту: 210 − 30 − 15 = 165 мм
TEXT_WIDTH = Mm(165)


def _set_font(style_or_run, name: str = FONT, size: float = 14, bold=None, italic=None):
    """Встановлює шрифт (включно з кириличною та східноазійською гарнітурою)."""
    font = style_or_run.font
    font.name = name
    font.size = Pt(size)
    font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    # Word зберігає окремі шрифти для різних наборів символів — задаємо всі
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(attr), None)


def _add_field(paragraph, instr: str, placeholder: str = ""):
    """Вставляє поле Word (номер сторінки, зміст тощо) у абзац."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    text_run = paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)
    return text_run


class DstuReport:
    """Обгортка над python-docx з методами для елементів звіту за ДСТУ."""

    def __init__(self, template_path):
        self.doc = Document(str(template_path))
        self.section_no = 0      # номер поточного розділу
        self.figure_no = 0       # лічильник рисунків у розділі
        self.table_no = 0        # лічильник таблиць у розділі
        self.formula_no = 0      # лічильник формул у розділі
        self.appendix = None     # літера поточного додатка (А, Б, …)
        self.counts = {"figures": 0, "tables": 0, "formulas": 0}
        self.headings = []       # (рівень, текст) усіх заголовків — для змісту
        self._setup_page()
        self._setup_styles()
        self._prepare_title_page()

    # ------------------------------------------------------------------
    # Налаштування сторінки, стилів і титульної сторінки
    # ------------------------------------------------------------------
    def _setup_page(self):
        """А4, поля за ДСТУ, номер сторінки праворуч угорі (крім титульної)."""
        for s in self.doc.sections:
            s.page_width, s.page_height = Mm(210), Mm(297)
            s.left_margin, s.right_margin = Mm(30), Mm(15)
            s.top_margin, s.bottom_margin = Mm(20), Mm(20)
            s.header_distance = Mm(10)
            s.different_first_page_header_footer = True  # титульна — без номера
            p = s.header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Cm(0)
            r = _add_field(p, "PAGE", "2")
            _set_font(r, size=14)

    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        _set_font(normal, size=14)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

        # Заголовки розділів і підрозділів: напівжирні, без кольору
        for name, size in (("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 14)):
            h = st[name]
            _set_font(h, size=size, bold=True, italic=False)
            hpf = h.paragraph_format
            hpf.line_spacing = 1.5
            hpf.first_line_indent = Cm(1.25)
            hpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            hpf.space_before = Pt(12) if name != "Heading 1" else Pt(0)
            hpf.space_after = Pt(12)
            hpf.keep_with_next = True

        cap = st["Caption"]
        _set_font(cap, size=14, bold=False, italic=False)
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.line_spacing = 1.0

        for name in ("Header", "TOC Heading"):
            if name in [s.name for s in st]:
                _set_font(st[name], size=14)

        # Стиль для лістингу коду (додаток): моноширинний 9 пт, одинарний інтервал
        from docx.enum.style import WD_STYLE_TYPE
        code = st.add_style("Code Listing", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        _set_font(code, name=CODE_FONT, size=9)
        cpf = code.paragraph_format
        cpf.line_spacing = 1.0
        cpf.first_line_indent = Cm(0)
        cpf.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Стилі елементів змісту: TOC 1 / TOC 2 — шрифт 14, без відступу першого рядка
        for lvl, indent in ((1, 0), (2, 0.5)):
            name = f"toc {lvl}"
            try:
                t = st[name]
            except KeyError:
                t = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                t.base_style = normal
            _set_font(t, size=14)
            t.paragraph_format.first_line_indent = Cm(0)
            t.paragraph_format.left_indent = Cm(indent)
            t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _prepare_title_page(self):
        """Прибирає абзацні відступи на титульній сторінці та зайвий порожній абзац."""
        body = self.doc.element.body
        paragraphs = self.doc.paragraphs
        for p in paragraphs:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                _set_font(r, size=14, bold=r.bold)
        # Порожні абзаци в кінці титульної сторінки (зокрема з розривом сторінки)
        # видаляємо: кожна наступна структурна частина сама починається з нової сторінки
        for p in reversed(paragraphs):
            if p.text.strip():
                break
            body.remove(p._p)

    # ------------------------------------------------------------------
    # Елементи тексту
    # ------------------------------------------------------------------
    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def structural_heading(self, text: str, in_toc: bool = True):
        """Заголовок структурної частини (ВСТУП, ВИСНОВКИ…): великими, по центру."""
        self.page_break()
        if in_toc:
            p = self.doc.add_paragraph(text.upper(), style="Heading 1")
            self.headings.append((1, text.upper()))
        else:
            p = self.doc.add_paragraph()
            r = p.add_run(text.upper())
            _set_font(r, bold=True)
            p.paragraph_format.space_after = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        return p

    def section(self, title: str):
        """Розділ з нової сторінки: «1 НАЗВА РОЗДІЛУ»."""
        self.section_no += 1
        self.figure_no = self.table_no = self.formula_no = 0
        self.page_break()
        text = f"{self.section_no} {title.upper()}"
        self.headings.append((1, text))
        return self.doc.add_paragraph(text, style="Heading 1")

    def subsection(self, number: str, title: str):
        """Підрозділ: «1.1 Назва підрозділу»."""
        self.headings.append((2, f"{number} {title}"))
        return self.doc.add_paragraph(f"{number} {title}", style="Heading 2")

    def para(self, text: str = "", bold_prefix: str = None, indent: bool = True, align=None):
        """Звичайний абзац; bold_prefix — напівжирний початок (напр. «Висновок.»)."""
        p = self.doc.add_paragraph()
        if bold_prefix:
            r = p.add_run(bold_prefix + " ")
            r.bold = True
        self._add_rich(p, text)
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        if align is not None:
            p.alignment = align
        return p

    @staticmethod
    def _add_rich(p, text: str):
        """Додає текст; фрагменти між символами ` виводяться моноширинним шрифтом."""
        parts = text.split("`")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            if i % 2 == 1:
                _set_font(r, name=CODE_FONT, size=12)

    def bullets(self, items: list, numbered: bool = False):
        """Перелік за ДСТУ: кожен елемент з абзацного відступу, через дефіс
        або з номером «1)»; у кінці елементів — крапка з комою / крапка."""
        for i, item in enumerate(items, 1):
            marker = f"{i})" if numbered else "–"
            p = self.doc.add_paragraph()
            self._add_rich(p, f"{marker} {item}")

    def formula(self, text: str):
        """Формула по центру з номером «(розділ.номер)» біля правого поля."""
        self.formula_no += 1
        self.counts["formulas"] += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        tabs = p.paragraph_format.tab_stops
        tabs.add_tab_stop(TEXT_WIDTH // 2, WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t" + text)
        r.italic = True
        p.add_run(f"\t({self.section_no}.{self.formula_no})")
        return f"{self.section_no}.{self.formula_no}"

    # ------------------------------------------------------------------
    # Рисунки і таблиці
    # ------------------------------------------------------------------
    def _number(self, kind: str) -> str:
        prefix = self.appendix if self.appendix else str(self.section_no)
        if kind == "figure":
            self.figure_no += 1
            self.counts["figures"] += 1
            return f"{prefix}.{self.figure_no}"
        self.table_no += 1
        self.counts["tables"] += 1
        return f"{prefix}.{self.table_no}"

    def next_figure_number(self) -> str:
        """Номер, який отримає наступний рисунок (для посилань у тексті до нього)."""
        prefix = self.appendix if self.appendix else str(self.section_no)
        return f"{prefix}.{self.figure_no + 1}"

    def next_table_number(self) -> str:
        prefix = self.appendix if self.appendix else str(self.section_no)
        return f"{prefix}.{self.table_no + 1}"

    def figure(self, image_path, caption: str, width_mm: float = 160):
        """Рисунок по центру, під ним підпис «Рисунок N.M – Назва»."""
        num = self._number("figure")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(image_path), width=Mm(min(width_mm, 165)))
        cap = self.doc.add_paragraph(f"Рисунок {num} – {caption}", style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(12)
        return num

    def table(self, header: list, rows: list, caption: str, col_widths_mm: list = None,
              font_size: float = 12, align_numbers: bool = True):
        """Таблиця з назвою над нею: «Таблиця N.M – Назва».

        Шрифт у таблиці — 12 пт (ДСТУ дозволяє менший кегль, ніж в основному тексті).
        """
        num = self._number("table")
        cap = self.doc.add_paragraph(f"Таблиця {num} – {caption}")
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT

        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if col_widths_mm is None:
            col_widths_mm = [165 / len(header)] * len(header)
        # Фіксована розмітка: ширини стовпців задаються явно і в сітці таблиці
        # (tblGrid — її використовують LibreOffice/Google Docs), і в кожній
        # клітинці (її використовує Word)
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, w in enumerate(col_widths_mm):
            grid.gridCol_lst[i].w = Mm(w)
            for cell in t.columns[i].cells:
                cell.width = Mm(w)

        def fill(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = align
            r = p.add_run(str(text))
            _set_font(r, size=font_size, bold=bold)

        for j, h in enumerate(header):
            fill(t.rows[0].cells[j], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        # Рядок заголовка повторюється при перенесенні таблиці на наступну сторінку
        trpr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trpr.append(hdr)

        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                is_num = align_numbers and _looks_numeric(val)
                fill(t.rows[i].cells[j], val,
                     align=WD_ALIGN_PARAGRAPH.CENTER if is_num else WD_ALIGN_PARAGRAPH.LEFT)
        # Невеликий відступ після таблиці
        after = self.doc.add_paragraph()
        after.paragraph_format.line_spacing = 1.0
        return num

    # ------------------------------------------------------------------
    # Зміст, лістинги, додатки
    # ------------------------------------------------------------------
    def toc(self, entries: list = None):
        """Сторінка ЗМІСТ.

        Зміст вставляється як поле Word TOC, але з уже заповненим результатом:
        entries — список (рівень, текст, номер сторінки), отриманий після
        першого проходу формування звіту (див. build_report.py). Тому зміст
        видно одразу, а за потреби Word оновить його клавішею F9.
        """
        self.structural_heading("Зміст", in_toc=False)
        entries = entries or [(1, "Зміст буде сформовано", "")]
        paragraphs = []
        for level, text, page in entries:
            p = self.doc.add_paragraph(style=f"toc {level}")
            p.paragraph_format.tab_stops.add_tab_stop(
                TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            paragraphs.append((p, text, page))
        # Початок поля — у першому абзаці, кінець — в останньому
        first = paragraphs[0][0]
        run = first.add_run()
        for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                                 ("w:instrText", {"xml:space": "preserve"}, 'TOC \\o "1-2" \\h \\z \\u'),
                                 ("w:fldChar", {"w:fldCharType": "separate"}, None)):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            if text:
                el.text = text
            run._r.append(el)
        for p, text, page in paragraphs:
            p.add_run(f"{text}\t{page}")
        end = paragraphs[-1][0].add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        end._r.append(fld_end)

    def appendix_heading(self, letter: str, title: str):
        """Додаток з нової сторінки: «ДОДАТОК А» і назва по центру."""
        self.appendix = letter
        self.figure_no = self.table_no = 0
        self.page_break()
        p = self.doc.add_paragraph(f"ДОДАТОК {letter}", style="Heading 1")
        self.headings.append((1, f"ДОДАТОК {letter} {title.upper()}"))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        t = self.doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.first_line_indent = Cm(0)
        t.paragraph_format.space_after = Pt(12)
        r = t.add_run(title)
        r.bold = True

    def code_listing(self, title: str, code: str):
        """Лістинг файлу з кодом: заголовок і текст моноширинним шрифтом."""
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.bold = True
        for line in code.rstrip("\n").split("\n"):
            cp = self.doc.add_paragraph(style="Code Listing")
            # Зберігаємо відступи коду: пробіли на початку рядка не стискаються
            cp.add_run(line if line else " ")

    def save(self, path):
        self.doc.save(str(path))


def _looks_numeric(val) -> bool:
    """Чи схоже значення клітинки на число (тоді вирівнюємо по центру)."""
    s = str(val).replace(",", ".").replace("%", "").replace(" ", "").replace("−", "-")
    s = s.replace("×", "").replace("/", "").replace("–", "")
    try:
        float(s)
        return True
    except ValueError:
        return False
